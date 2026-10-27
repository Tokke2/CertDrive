from __future__ import annotations

"""
plugins/certificates/__init__.py
Intygsgenerator — paketgränssnitt för intyg, mallar, dependency-status och
robust degraded mode.

Paketet exporterar alltid ett användbart API. Den verkliga mallskanningen ägs
av plugin.py och den cachade DOCX-/PDF-motorn av engine.py.
"""

import importlib.util
import json
import logging
import os
import re
import subprocess
import sys
import threading
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from PySide6.QtWidgets import QWidget

__version__: str = "10.0.0"
__author__: str = "Rickard Längkvist"
__all__: list[str]

_START_YEAR: int = 2025
_CURRENT_YEAR: int = datetime.now().year
COPYRIGHT: str = (
    f"© Rickard Längkvist {_START_YEAR}–{_CURRENT_YEAR}"
    if _CURRENT_YEAR > _START_YEAR
    else f"© Rickard Längkvist {_START_YEAR}"
)

PLUGIN_NAME: str = "Intyg"
PLUGIN_VERSION: str = __version__
PLUGIN_AUTHOR: str = __author__
PLUGIN_PRIORITY: int = 100
FOLDER_CONFIG: str = "js"

_logger: logging.Logger = logging.getLogger(__name__)
_state_lock = threading.RLock()
_pdf_converter_lock = threading.RLock()
_dependency_lock = threading.RLock()

_ENGINE_LOADED: bool = False
_PLUGIN_LOADED: bool = False
_ENGINE_ERROR: str = ""
_PLUGIN_ERROR: str = ""
_pdf_converter_cache: Any | None = None
_dependency_status_cache: dict[str, bool] | None = None
_engine_get_field_label: Any | None = None
_engine_is_auto_field: Any | None = None

PLACEHOLDER_PATTERN: re.Pattern[str] = re.compile(
    r"::\s*([A-Za-z0-9_åäöÅÄÖ]+)\s*::"
)
VIN_PATTERN: re.Pattern[str] = re.compile(
    r"^[A-HJ-NPR-Z0-9]{17}$",
    re.IGNORECASE,
)
REG_PATTERN: re.Pattern[str] = re.compile(
    r"^[A-ZÅÄÖ]{3}\s?\d{2,3}\s?[A-ZÅÄÖ0-9]?$",
    re.IGNORECASE,
)

_AUTO_NAMES: frozenset[str] = frozenset(
    {
        "intyg",
        "nummer",
        "number",
        "nr",
        "löpnummer",
        "id",
        "datum",
        "date",
        "idag",
        "today",
        "år_nu",
        "currentyear",
        "månad",
        "month",
        "dag",
        "day",
        "tid",
        "time",
        "timestamp",
    }
)
_NEVER_PREFILL: frozenset[str] = frozenset(
    {"år", "year", "årsmodell", "modelyear"}
)

FIELD_LABELS: dict[str, str] = {
    "vin": "Chassinummer (VIN)",
    "chassi": "Chassinummer",
    "reg": "Registreringsnummer",
    "regnr": "Registreringsnummer",
    "registreringsnummer": "Registreringsnummer",
    "märke": "Fordonsmärke",
    "marke": "Fordonsmärke",
    "modell": "Fordonsmodell",
    "år": "Årsmodell",
    "year": "Årsmodell",
    "årsmodell": "Årsmodell",
    "namn": "Namn",
    "ägare": "Ägare",
    "agare": "Ägare",
    "kund": "Kund",
    "adress": "Adress",
    "telefon": "Telefon",
    "email": "E-post",
    "e-post": "E-post",
    "anteckning": "Anteckning",
    "kommentar": "Kommentar",
}
DEFAULT_VALUES: dict[str, str] = {"land": "Sverige"}


class _AutoFieldSet:
    """Case-insensitiv uppsättning för automatiska placeholders."""

    def __contains__(self, item: object) -> bool:
        return isinstance(item, str) and item.casefold() in _AUTO_NAMES

    def __iter__(self):
        yield from sorted(_AUTO_NAMES)

    def __len__(self) -> int:
        return len(_AUTO_NAMES)

    def __repr__(self) -> str:
        return f"AutoFieldSet({len(self)} fields)"


AUTO_FIELDS: _AutoFieldSet = _AutoFieldSet()


class ConversionMethod(Enum):
    """Fallback för engine.ConversionMethod."""

    NONE = 0
    LIBREOFFICE = 1
    DOCX2PDF = 2
    COMTYPES = 3
    UNOCONV = 4
    ABIWORD = 5

    @property
    def display_name(self) -> str:
        return {
            ConversionMethod.NONE: "Ingen",
            ConversionMethod.LIBREOFFICE: "LibreOffice",
            ConversionMethod.DOCX2PDF: "docx2pdf",
            ConversionMethod.COMTYPES: "Word COM",
            ConversionMethod.UNOCONV: "unoconv",
            ConversionMethod.ABIWORD: "AbiWord",
        }[self]

    @property
    def is_available(self) -> bool:
        return self is not ConversionMethod.NONE


class ValidationResult(Enum):
    """Fallback för engine.ValidationResult."""

    VALID = 1
    WARNING = 2
    ERROR = 3


class ScanStatus(Enum):
    """Fallback för engine.ScanStatus."""

    IDLE = 1
    SCANNING = 2
    COMPLETE = 3
    ERROR = 4


def _as_bool(value: Any, default: bool) -> bool:
    """Normalisera bool-värden från JSON-konfiguration."""
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        lowered = value.strip().casefold()
        if lowered in {"true", "1", "yes", "ja"}:
            return True
        if lowered in {"false", "0", "no", "nej"}:
            return False
    return default


def _path(value: Any) -> Path:
    """Konvertera till Path utan att låta trasig mallmetadata stoppa appen."""
    try:
        return Path(value)
    except (TypeError, ValueError):
        return Path()


@dataclass(slots=True)
class TemplateConfig:
    """Konfiguration för en mall när engine inte kan laddas."""

    name: str = ""
    description: str = ""
    category: str = ""
    counter_prefix: str = ""
    counter_start: int = 1
    extra_fields: list[str] = field(default_factory=list)
    default_values: dict[str, str] = field(default_factory=dict)
    pdf_attachments_order: list[str] = field(default_factory=list)
    validate_vin: bool = True
    validate_reg: bool = False
    output_format: str = "pdf"
    filename_pattern: str = "{number} Intyg {prefix} {vin} {reg} {date}"

    @classmethod
    def from_dict(
        cls,
        data: dict[str, Any],
        name: str = "",
    ) -> TemplateConfig:
        """Skapa en säker config från mallens valfria JSON-fil."""
        if not isinstance(data, dict):
            return cls(name=name)

        raw_start = data.get("counter_start", 1)
        try:
            counter_start = max(1, int(raw_start))
        except (TypeError, ValueError):
            counter_start = 1

        raw_defaults = data.get("default_values", data.get("defaults", {}))
        defaults: dict[str, str] = {}
        if isinstance(raw_defaults, dict):
            defaults = {
                str(key): str(value)
                for key, value in raw_defaults.items()
                if str(key).casefold() not in _NEVER_PREFILL
            }

        raw_fields = data.get("extra_fields", data.get("fields", []))
        extra_fields = (
            [str(value) for value in raw_fields if str(value).strip()]
            if isinstance(raw_fields, list)
            else []
        )
        raw_order = data.get(
            "pdf_attachments_order",
            data.get("attachments", []),
        )
        attachments = (
            [str(value) for value in raw_order if str(value).strip()]
            if isinstance(raw_order, list)
            else []
        )

        return cls(
            name=str(data.get("name", name) or name),
            description=str(data.get("description", "") or ""),
            category=str(data.get("category", "") or ""),
            counter_prefix=str(data.get("counter_prefix", "") or ""),
            counter_start=counter_start,
            extra_fields=list(dict.fromkeys(extra_fields)),
            default_values=defaults,
            pdf_attachments_order=list(dict.fromkeys(attachments)),
            validate_vin=_as_bool(data.get("validate_vin"), True),
            validate_reg=_as_bool(data.get("validate_reg"), False),
            output_format=str(data.get("output_format", "pdf") or "pdf"),
            filename_pattern=str(
                data.get(
                    "filename_pattern",
                    "{number} Intyg {prefix} {vin} {reg} {date}",
                )
                or "{number} Intyg {prefix} {vin} {reg} {date}"
            ),
        )


@dataclass(slots=True)
class Template:
    """Mallmodell för degraded mode."""

    name: str = ""
    folder: Path = field(default_factory=Path)
    docx_file: Path = field(default_factory=Path)
    pdf_attachments: list[Path] = field(default_factory=list)
    config: TemplateConfig = field(default_factory=TemplateConfig)
    placeholders: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.name = str(self.name or "").strip()
        self.folder = _path(self.folder)
        self.docx_file = _path(self.docx_file)
        self.pdf_attachments = [
            _path(path) for path in self.pdf_attachments if _path(path).suffix.casefold() == ".pdf"
        ]
        if not self.config.name:
            self.config.name = self.name
        self.placeholders = list(
            dict.fromkeys(str(value).strip() for value in self.placeholders if str(value).strip())
        )

    @property
    def is_valid(self) -> bool:
        return (
            bool(self.name)
            and self.docx_file.is_file()
            and self.docx_file.suffix.casefold() == ".docx"
            and not self.docx_file.name.startswith("~$")
        )

    @property
    def attachment_count(self) -> int:
        return len(self.pdf_attachments)

    @property
    def has_attachments(self) -> bool:
        return bool(self.pdf_attachments)

    def get_user_fields(self) -> list[str]:
        fields = list(dict.fromkeys(self.placeholders + self.config.extra_fields))
        return [field_name for field_name in fields if field_name not in AUTO_FIELDS]

    def get_default(self, field_name: str) -> str:
        normalized = field_name.casefold()
        if normalized in _NEVER_PREFILL:
            return ""
        for key, value in self.config.default_values.items():
            if key.casefold() == normalized:
                return value
        return DEFAULT_VALUES.get(normalized, "")

    def get_ordered_attachments(self) -> list[Path]:
        available = [path for path in self.pdf_attachments if path.is_file()]
        if not self.config.pdf_attachments_order:
            return sorted(available, key=lambda path: path.name.casefold())

        order = [item.casefold() for item in self.config.pdf_attachments_order]
        priority = {
            item: position
            for position, item in enumerate(order)
        }
        return sorted(
            available,
            key=lambda path: (
                priority.get(path.name.casefold(), priority.get(path.stem.casefold(), len(priority))),
                path.name.casefold(),
            ),
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "folder": str(self.folder),
            "docx_file": str(self.docx_file),
            "pdf_attachments": [str(path) for path in self.pdf_attachments],
            "placeholders": list(self.placeholders),
            "user_fields": self.get_user_fields(),
            "is_valid": self.is_valid,
            "attachment_count": self.attachment_count,
            "config": {
                "name": self.config.name,
                "description": self.config.description,
                "category": self.config.category,
                "counter_prefix": self.config.counter_prefix,
                "counter_start": self.config.counter_start,
                "extra_fields": list(self.config.extra_fields),
                "default_values": dict(self.config.default_values),
                "pdf_attachments_order": list(self.config.pdf_attachments_order),
                "validate_vin": self.config.validate_vin,
                "validate_reg": self.config.validate_reg,
                "output_format": self.config.output_format,
                "filename_pattern": self.config.filename_pattern,
            },
        }


@dataclass(slots=True)
class GenerationResult:
    """Resultat från en intygsgenerering."""

    success: bool = False
    output_path: Path | None = None
    number: int = 0
    filename: str = ""
    page_count: int = 0
    error: str | None = None
    duration: float = 0.0
    warnings: list[str] = field(default_factory=list)
    template_name: str = ""

    @property
    def is_error(self) -> bool:
        return not self.success

    def to_dict(self) -> dict[str, Any]:
        return {
            "success": self.success,
            "output_path": str(self.output_path) if self.output_path else None,
            "number": self.number,
            "filename": self.filename,
            "page_count": self.page_count,
            "error": self.error,
            "duration": round(self.duration, 3),
            "warnings": list(self.warnings),
            "template_name": self.template_name,
        }


@dataclass(slots=True)
class FieldValidation:
    """Valideringsresultat för ett fält."""

    field_name: str = ""
    result: ValidationResult = ValidationResult.ERROR
    message: str = ""
    formatted_value: str = ""

    @property
    def is_valid(self) -> bool:
        return self.result in {ValidationResult.VALID, ValidationResult.WARNING}

    @property
    def is_error(self) -> bool:
        return self.result is ValidationResult.ERROR

    @property
    def is_warning(self) -> bool:
        return self.result is ValidationResult.WARNING


@dataclass(slots=True)
class BatchResult:
    """Resultat från batchgenerering."""

    total: int = 0
    succeeded: int = 0
    failed: int = 0
    results: list[GenerationResult] = field(default_factory=list)
    duration: float = 0.0

    @property
    def success(self) -> bool:
        return self.total > 0 and self.failed == 0 and self.succeeded == self.total

    def to_dict(self) -> dict[str, Any]:
        return {
            "total": self.total,
            "succeeded": self.succeeded,
            "failed": self.failed,
            "duration": round(self.duration, 3),
            "results": [result.to_dict() for result in self.results],
        }


class PlatformHelper:
    """Minimal, icke-blockerande plattformshjälpare för degraded mode."""

    @staticmethod
    def is_windows() -> bool:
        return sys.platform == "win32"

    @staticmethod
    def is_macos() -> bool:
        return sys.platform == "darwin"

    @staticmethod
    def is_linux() -> bool:
        return sys.platform.startswith("linux")

    @staticmethod
    def get_platform_name() -> str:
        if PlatformHelper.is_windows():
            return "Windows"
        if PlatformHelper.is_macos():
            return "macOS"
        if PlatformHelper.is_linux():
            return "Linux"
        return sys.platform

    @staticmethod
    def open_file(path: Path) -> bool:
        target = _path(path)
        if not target.exists():
            _logger.warning("Kan inte öppna saknad sökväg: %s", target)
            return False
        try:
            if PlatformHelper.is_windows():
                startfile = getattr(os, "startfile", None)
                if callable(startfile):
                    startfile(str(target))
                    return True
                return False
            command = ["open", str(target)] if PlatformHelper.is_macos() else ["xdg-open", str(target)]
            subprocess.Popen(
                command,
                start_new_session=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            return True
        except (OSError, ValueError) as exc:
            _logger.error("Kunde inte öppna %s: %s", target, exc)
            return False

    @staticmethod
    def open_folder(path: Path) -> bool:
        target = _path(path)
        if target.is_file():
            target = target.parent
        try:
            target.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            _logger.error("Kunde inte skapa mapp %s: %s", target, exc)
            return False
        return PlatformHelper.open_file(target)


class VINValidator:
    """Fallbackvalidering för VIN."""

    @staticmethod
    def validate(vin: str) -> FieldValidation:
        cleaned = re.sub(r"[^A-Z0-9]", "", str(vin).upper())
        if not cleaned:
            return FieldValidation("VIN", ValidationResult.ERROR, "VIN får inte vara tomt")
        if len(cleaned) != 17:
            return FieldValidation(
                "VIN",
                ValidationResult.ERROR,
                f"VIN måste vara 17 tecken (angivet: {len(cleaned)})",
                cleaned,
            )
        if any(character in cleaned for character in "IOQ"):
            return FieldValidation(
                "VIN",
                ValidationResult.ERROR,
                "VIN får inte innehålla I, O eller Q",
                cleaned,
            )
        return FieldValidation("VIN", ValidationResult.VALID, "Giltigt VIN", cleaned)

    @staticmethod
    def format_vin(vin: str) -> str:
        return re.sub(r"[^A-Z0-9]", "", str(vin).upper())[:17]


class RegValidator:
    """Fallbackvalidering för svenska registreringsnummer."""

    @staticmethod
    def validate(reg: str) -> FieldValidation:
        cleaned = re.sub(r"[^A-Z0-9ÅÄÖ]", "", str(reg).upper())
        if not cleaned:
            return FieldValidation("REG", ValidationResult.ERROR, "Registreringsnummer saknas")
        if not REG_PATTERN.fullmatch(cleaned):
            return FieldValidation(
                "REG",
                ValidationResult.WARNING,
                f"'{cleaned}' har ett ovanligt format",
                cleaned,
            )
        return FieldValidation("REG", ValidationResult.VALID, "Giltigt registreringsnummer", cleaned)


class PlaceholderParser:
    """DOCX-placeholderparser med cache per filstämpel.

    Den hanterar body, nästlade tabeller samt sidhuvud/-sidfot inklusive
    first/even-page-varianter och sammanfogar Word-runs innan regex-matchning.
    """

    _cache_lock = threading.RLock()
    _cache: dict[tuple[Path, int, int], tuple[str, ...]] = {}
    _cache_limit = 256

    @classmethod
    def clear_cache(cls) -> None:
        with cls._cache_lock:
            cls._cache.clear()

    @classmethod
    def find_placeholders(cls, docx_path: Path | str) -> list[str]:
        path = _path(docx_path)
        try:
            resolved = path.resolve()
            stat = resolved.stat()
        except OSError:
            return []
        if not resolved.is_file() or resolved.suffix.casefold() != ".docx":
            return []

        key = (resolved, stat.st_mtime_ns, stat.st_size)
        with cls._cache_lock:
            cached = cls._cache.get(key)
        if cached is not None:
            return list(cached)

        values = cls._scan_docx(resolved)
        with cls._cache_lock:
            if len(cls._cache) >= cls._cache_limit:
                cls._cache.pop(next(iter(cls._cache)))
            cls._cache[key] = tuple(values)
        return values

    @staticmethod
    def _scan_docx(path: Path) -> list[str]:
        try:
            from docx import Document
        except ImportError:
            _logger.warning("python-docx saknas; placeholders kan inte läsas")
            return []

        found: set[str] = set()
        seen_containers: set[int] = set()

        def scan_paragraphs(paragraphs: Any) -> None:
            for paragraph in paragraphs:
                runs = getattr(paragraph, "runs", [])
                text = "".join(str(getattr(run, "text", "") or "") for run in runs)
                if not text:
                    text = str(getattr(paragraph, "text", "") or "")
                for match in PLACEHOLDER_PATTERN.finditer(text):
                    found.add(match.group(1).strip())

        def scan_container(container: Any) -> None:
            if container is None or id(container) in seen_containers:
                return
            seen_containers.add(id(container))
            scan_paragraphs(getattr(container, "paragraphs", []))
            for table in getattr(container, "tables", []):
                for row in getattr(table, "rows", []):
                    for cell in getattr(row, "cells", []):
                        scan_container(cell)

        try:
            document = Document(str(path))
            scan_container(document)
            for section in document.sections:
                for container in (
                    section.header,
                    section.footer,
                    section.first_page_header,
                    section.first_page_footer,
                    section.even_page_header,
                    section.even_page_footer,
                ):
                    scan_container(container)
        except Exception:
            _logger.exception("Kunde inte skanna placeholders i %s", path)
            return []
        return sorted(found, key=str.casefold)

    @staticmethod
    def replace_in_docx(
        source: Path | str,
        destination: Path | str,
        replacements: dict[str, Any],
    ) -> bool:
        """Byt placeholders i DOCX i degraded mode utan tyst dataförlust."""
        try:
            from docx import Document
        except ImportError:
            _logger.error("python-docx saknas; DOCX kan inte fyllas")
            return False

        source_path = _path(source)
        destination_path = _path(destination)
        if not source_path.is_file():
            _logger.error("Källmall saknas: %s", source_path)
            return False

        lookup = {str(key).casefold(): str(value) for key, value in replacements.items()}

        def replace_paragraph(paragraph: Any) -> None:
            runs = list(getattr(paragraph, "runs", []))
            original = "".join(str(getattr(run, "text", "") or "") for run in runs)
            if not original:
                return

            def resolve(match: re.Match[str]) -> str:
                key = match.group(1).strip().casefold()
                return lookup.get(key, match.group(0))

            updated = PLACEHOLDER_PATTERN.sub(resolve, original)
            if updated == original:
                return
            if not runs:
                paragraph.add_run(updated)
                return
            runs[0].text = updated
            for run in runs[1:]:
                run.text = ""

        def replace_container(container: Any, seen: set[int]) -> None:
            if container is None or id(container) in seen:
                return
            seen.add(id(container))
            for paragraph in getattr(container, "paragraphs", []):
                replace_paragraph(paragraph)
            for table in getattr(container, "tables", []):
                for row in getattr(table, "rows", []):
                    for cell in getattr(row, "cells", []):
                        replace_container(cell, seen)

        try:
            document = Document(str(source_path))
            seen_containers: set[int] = set()
            replace_container(document, seen_containers)
            for section in document.sections:
                for container in (
                    section.header,
                    section.footer,
                    section.first_page_header,
                    section.first_page_footer,
                    section.even_page_header,
                    section.even_page_footer,
                ):
                    replace_container(container, seen_containers)
            destination_path.parent.mkdir(parents=True, exist_ok=True)
            document.save(str(destination_path))
            return True
        except Exception:
            _logger.exception("Kunde inte ersätta placeholders i %s", source_path)
            return False


class PDFConverter:
    """Tydlig degraded-mode-konverterare när engine saknas."""

    @property
    def is_available(self) -> bool:
        return False

    @property
    def method_name(self) -> str:
        return "Ingen (engine ej laddad)"

    def refresh(self) -> None:
        """Behåll kompatibelt API; tillgänglighet ändras inte utan engine."""

    def convert(self, *_args: Any, **_kwargs: Any) -> Path:
        raise RuntimeError("DOCX→PDF kräver plugins/certificates/engine.py")


class PDFMerger:
    """Fallback för PDF-sammanslagning med pypdf när det finns installerat."""

    @staticmethod
    def merge(files: list[Path], destination: Path, log: Any = None) -> tuple[bool, int]:
        logger = log or _logger
        valid = [
            _path(path)
            for path in files
            if _path(path).is_file() and _path(path).suffix.casefold() == ".pdf"
        ]
        if not valid:
            logger.error("PDF-sammanslagning saknar giltiga indatafiler")
            return False, 0
        try:
            from pypdf import PdfReader, PdfWriter

            writer = PdfWriter()
            pages = 0
            for file_path in valid:
                reader = PdfReader(str(file_path))
                for page in reader.pages:
                    writer.add_page(page)
                    pages += 1
            destination_path = _path(destination)
            destination_path.parent.mkdir(parents=True, exist_ok=True)
            temporary = destination_path.with_suffix(f"{destination_path.suffix}.tmp")
            with temporary.open("wb") as handle:
                writer.write(handle)
            temporary.replace(destination_path)
            return True, pages
        except Exception:
            logger.exception("Kunde inte slå ihop PDF-filer")
            return False, 0

    @staticmethod
    def get_page_count(path: Path) -> int:
        try:
            from pypdf import PdfReader

            return len(PdfReader(str(_path(path))).pages)
        except Exception:
            _logger.debug("Kunde inte räkna PDF-sidor för %s", path, exc_info=True)
            return 0


class CounterManager:
    """Trådsäker, persistent räknare för degraded mode."""

    _CONFIG_KEY = "certificates.counters"
    _DEFAULT_FILE = "certificates_counters.json"

    def __init__(
        self,
        config_manager: Any | None = None,
        config_path: Path | str | None = None,
        **_kwargs: Any,
    ) -> None:
        self._lock = threading.RLock()
        self._config_manager = config_manager
        default_path = Path(__file__).resolve().parent.parent.parent / FOLDER_CONFIG
        self._path = _path(config_path) if config_path is not None else default_path / self._DEFAULT_FILE
        self._values: dict[str, int] = {}
        self._load()

    def _load(self) -> None:
        raw: Any = None
        getter = getattr(self._config_manager, "get", None)
        if callable(getter):
            try:
                raw = getter(self._CONFIG_KEY, {})
            except Exception:
                _logger.exception("Kunde inte läsa räknare från ConfigManager")
        if (not isinstance(raw, dict) or not raw) and self._path.is_file():
            try:
                raw = json.loads(self._path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                _logger.exception("Kunde inte läsa räknarfil %s", self._path)
        if not isinstance(raw, dict):
            return
        values = raw.get("counters", raw)
        if not isinstance(values, dict):
            return
        with self._lock:
            self._values = {
                str(key): max(0, int(value))
                for key, value in values.items()
                if self._is_integer(value)
            }

    @staticmethod
    def _is_integer(value: Any) -> bool:
        try:
            int(value)
            return not isinstance(value, bool)
        except (TypeError, ValueError):
            return False

    def get_current(self, name: str) -> int:
        with self._lock:
            return self._values.get(str(name), 0)

    def get_next(self, name: str, start: int = 1) -> int:
        try:
            minimum = max(1, int(start))
        except (TypeError, ValueError):
            minimum = 1
        with self._lock:
            return max(self._values.get(str(name), 0) + 1, minimum)

    def increment(self, name: str, start: int = 1) -> int:
        key = str(name)
        with self._lock:
            value = self.get_next(key, start)
            self._values[key] = value
            return value

    def set(self, name: str, value: int) -> None:
        if isinstance(value, bool):
            raise ValueError("Räknarvärde måste vara ett heltal")
        with self._lock:
            self._values[str(name)] = max(0, int(value))

    def get_all(self) -> dict[str, int]:
        with self._lock:
            return dict(self._values)

    def get_total(self) -> int:
        with self._lock:
            return sum(self._values.values())

    def save(self) -> None:
        """Spara räknare via ConfigManager eller atomisk JSON-skrivning."""
        values = self.get_all()
        setter = getattr(self._config_manager, "set", None)
        if callable(setter):
            try:
                setter(self._CONFIG_KEY, values)
                return
            except Exception:
                _logger.exception("Kunde inte spara räknare via ConfigManager")
        try:
            self._path.parent.mkdir(parents=True, exist_ok=True)
            temporary = self._path.with_suffix(f"{self._path.suffix}.tmp")
            temporary.write_text(
                json.dumps({"counters": values}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            temporary.replace(self._path)
        except OSError:
            _logger.exception("Kunde inte spara räknarfil %s", self._path)


class CertificateGenerator:
    """Säker generator som rapporterar tydligt varför engine krävs."""

    def __init__(self, *_args: Any, **_kwargs: Any) -> None:
        self._reason = "Intygsmotorn kunde inte laddas; se loggen för engine-felet."

    def generate(self, *_args: Any, **_kwargs: Any) -> GenerationResult:
        return GenerationResult(success=False, error=self._reason)

    def generate_batch(self, *_args: Any, **_kwargs: Any) -> BatchResult:
        batch_values = _args[1] if len(_args) > 1 and isinstance(_args[1], list) else []
        total = len(batch_values)
        return BatchResult(total=total, failed=total or 1)


def safe_filename(name: str, max_length: int = 200) -> str:
    """Skapa säkert filnamn och behåll svenska bokstäver via NFC."""
    normalized = unicodedata.normalize("NFC", str(name or "")).strip()
    if not normalized:
        return "namnlös"
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", normalized)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(". ")
    return (cleaned[:max(1, max_length)].rstrip(". ") or "namnlös")


# engine importeras efter att fallback-API:t är komplett. Om importen fallerar
# fortsätter paketet fungera för diagnostik och placeholder-läsning.
try:
    from .engine import (
        AUTO_FIELDS,
        COPYRIGHT,
        DEFAULT_VALUES,
        FIELD_LABELS,
        PLACEHOLDER_PATTERN,
        PLUGIN_AUTHOR,
        PLUGIN_NAME,
        PLUGIN_PRIORITY,
        PLUGIN_VERSION,
        REG_PATTERN,
        VIN_PATTERN,
        BatchResult,
        CertificateGenerator,
        ConversionMethod,
        CounterManager,
        FieldValidation,
        GenerationResult,
        PDFConverter,
        PDFMerger,
        PlaceholderParser,
        PlatformHelper,
        RegValidator,
        ScanStatus,
        Template,
        TemplateConfig,
        VINValidator,
        ValidationResult,
        get_field_label as _engine_get_field_label,
        is_auto_field as _engine_is_auto_field,
        safe_filename,
    )

    with _state_lock:
        _ENGINE_LOADED = True
        _ENGINE_ERROR = ""
except Exception as exc:
    with _state_lock:
        _ENGINE_LOADED = False
        _ENGINE_ERROR = f"{type(exc).__name__}: {exc}"
    _logger.critical(
        "[CERTIFICATES] engine.py kunde inte laddas; degraded mode aktiveras: %s",
        exc,
        exc_info=True,
    )

try:
    from .plugin import Plugin

    with _state_lock:
        _PLUGIN_LOADED = True
        _PLUGIN_ERROR = ""
except Exception as exc:
    with _state_lock:
        _PLUGIN_LOADED = False
        _PLUGIN_ERROR = f"{type(exc).__name__}: {exc}"
    _logger.critical(
        "[CERTIFICATES] plugin.py kunde inte laddas: %s",
        exc,
        exc_info=True,
    )

    class Plugin:
        """Fail-safe plugin som gör felet synligt utan att krascha övriga appen."""

        name: str = PLUGIN_NAME
        version: str = PLUGIN_VERSION
        description: str = "Intyg och mallar i degraded mode"
        priority: int = PLUGIN_PRIORITY

        def __init__(self) -> None:
            self._context: Any | None = None
            self._tab: QWidget | None = None

        def on_load(self, context: Any) -> None:
            self._context = context
            emitter = getattr(context, "emit_event", None)
            if callable(emitter):
                try:
                    emitter(
                        "error.occurred",
                        {
                            "source": "certificates",
                            "message": f"Intyg-plugin kunde inte laddas: {_PLUGIN_ERROR}",
                            "level": "error",
                        },
                    )
                except Exception:
                    _logger.exception("Kunde inte emittera certificates-fel")

        def setup(self, context: Any) -> None:
            self.on_load(context)

        def on_unload(self) -> None:
            self._context = None
            self._tab = None

        def on_close(self) -> None:
            self.on_unload()

        def get_tab(self) -> QWidget | None:
            if self._tab is not None:
                return self._tab
            try:
                from PySide6.QtCore import Qt
                from PySide6.QtWidgets import QLabel, QVBoxLayout, QWidget

                widget = QWidget()
                layout = QVBoxLayout(widget)
                layout.setContentsMargins(28, 28, 28, 28)
                layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
                title = QLabel("Intygsmotorn kunde inte laddas", widget)
                title_font = title.font()
                title_font.setBold(True)
                title_font.setPointSize(max(14, title_font.pointSize() + 3))
                title.setFont(title_font)
                title.setAlignment(Qt.AlignmentFlag.AlignCenter)
                layout.addWidget(title)
                detail = QLabel(
                    "Mallar kan inte visas förrän plugin.py/engine.py har laddats.\n\n"
                    f"Detalj: {_PLUGIN_ERROR}",
                    widget,
                )
                detail.setWordWrap(True)
                detail.setAlignment(Qt.AlignmentFlag.AlignCenter)
                layout.addWidget(detail)
                self._tab = widget
                return widget
            except Exception:
                _logger.exception("Kunde inte bygga degraded certificates-flik")
                return None


def _check_dependencies() -> dict[str, bool]:
    """Kontrollera paket utan att importera eller blockera appstarten."""
    global _dependency_status_cache
    with _dependency_lock:
        if _dependency_status_cache is not None:
            return dict(_dependency_status_cache)

    status = {
        "python-docx": importlib.util.find_spec("docx") is not None,
        "pypdf": (
            importlib.util.find_spec("pypdf") is not None
            or importlib.util.find_spec("PyPDF2") is not None
        ),
        "docx2pdf": importlib.util.find_spec("docx2pdf") is not None,
    }
    with _dependency_lock:
        _dependency_status_cache = dict(status)
    return status


def get_pdf_converter() -> Any | None:
    """Returnera en lazily skapad engine-konverterare, aldrig vid paketimport."""
    global _pdf_converter_cache
    if not _ENGINE_LOADED:
        return None
    with _pdf_converter_lock:
        if _pdf_converter_cache is not None:
            return _pdf_converter_cache
        try:
            _pdf_converter_cache = PDFConverter(log=_logger)
        except TypeError:
            _pdf_converter_cache = PDFConverter()
        except Exception:
            _logger.exception("Kunde inte initiera PDFConverter")
            return None
        return _pdf_converter_cache


def get_plugin_info() -> dict[str, Any]:
    """Returnera aktuell metadata och reella importstatusar."""
    with _state_lock:
        engine_loaded = _ENGINE_LOADED
        plugin_loaded = _PLUGIN_LOADED
        engine_error = _ENGINE_ERROR
        plugin_error = _PLUGIN_ERROR
    return {
        "name": PLUGIN_NAME,
        "version": PLUGIN_VERSION,
        "author": PLUGIN_AUTHOR,
        "description": "Mallhantering och intygsgenerering med PDF-export",
        "copyright": COPYRIGHT,
        "priority": PLUGIN_PRIORITY,
        "tab_priority": PLUGIN_PRIORITY,
        "has_tab": True,
        "min_python": "3.12",
        "platforms": ["windows", "darwin", "linux"],
        "requires": [],
        "optional": [
            "python-docx==0.8.11",
            "pypdf==5.3.1",
            "docx2pdf==0.1.8",
        ],
        "engine_loaded": engine_loaded,
        "plugin_loaded": plugin_loaded,
        "fully_loaded": engine_loaded and plugin_loaded,
        "engine_error": engine_error,
        "plugin_error": plugin_error,
        "paths": {"templates": "Mallar", "output": "Out", "config": FOLDER_CONFIG},
        "events": {
            "emits": [
                "certificate.before_generate",
                "certificate.generated",
                "certificate.generation_failed",
                "certificate.template_selected",
                "certificate.placeholder_updated",
                "certificates.imported",
                "certificates.batch.started",
                "certificates.batch.completed",
                "templates.scanned",
                "statistics.certificate_generated",
                "statistics.updated",
                "tray.status_update",
            ],
            "listens_to": [
                "refresh_templates",
                "settings.changed",
                "settings.theme_changed",
                "template.created",
                "template.deleted",
                "template.updated",
                "statistics.request",
            ],
        },
    }


PLUGIN_INFO: dict[str, Any] = get_plugin_info()
__copyright__: str = COPYRIGHT
__description__: str = "Mallhantering och intygsgenerering med PDF-export"


def get_version() -> str:
    return PLUGIN_VERSION


def get_copyright() -> str:
    return COPYRIGHT


def is_fully_loaded() -> bool:
    with _state_lock:
        return _ENGINE_LOADED and _PLUGIN_LOADED


def is_engine_loaded() -> bool:
    with _state_lock:
        return _ENGINE_LOADED


def is_plugin_loaded() -> bool:
    with _state_lock:
        return _PLUGIN_LOADED


def get_load_errors() -> dict[str, str]:
    with _state_lock:
        return {
            key: value
            for key, value in {
                "engine": _ENGINE_ERROR,
                "plugin": _PLUGIN_ERROR,
            }.items()
            if value
        }


def validate_vin(vin: str) -> FieldValidation:
    return VINValidator.validate(vin)


def validate_reg(reg: str) -> FieldValidation:
    return RegValidator.validate(reg)


def format_vin(vin: str) -> str:
    formatter = getattr(VINValidator, "format_vin", None)
    if callable(formatter):
        return str(formatter(vin))
    return re.sub(r"[^A-Z0-9]", "", str(vin).upper())[:17]


def find_placeholders(docx_path: Path | str) -> list[str]:
    """Läs placeholders med engine-cache eller komplett fallback-parser."""
    return list(PlaceholderParser.find_placeholders(_path(docx_path)))


def is_auto_field(field_name: str) -> bool:
    if callable(_engine_is_auto_field):
        try:
            return bool(_engine_is_auto_field(field_name))
        except Exception:
            _logger.exception("Engine kunde inte kontrollera auto-fält %s", field_name)
    return field_name in AUTO_FIELDS


def get_field_label(field_name: str, default: str | None = None) -> str:
    if callable(_engine_get_field_label):
        try:
            result = _engine_get_field_label(field_name)
            if result:
                return str(result)
        except Exception:
            _logger.exception("Engine kunde inte hämta fältetikett %s", field_name)
    normalized = str(field_name).casefold()
    return FIELD_LABELS.get(normalized, default or str(field_name))


def is_pdf_converter_available() -> bool:
    converter = get_pdf_converter()
    if converter is None:
        return False
    try:
        return bool(getattr(converter, "is_available", False))
    except Exception:
        _logger.exception("Kunde inte kontrollera PDFConverter")
        return False


def get_pdf_converter_name() -> str:
    converter = get_pdf_converter()
    if converter is None:
        return "Ingen (engine ej laddad)" if not _ENGINE_LOADED else "Ingen"
    try:
        return str(getattr(converter, "method_name", "Ingen"))
    except Exception:
        _logger.exception("Kunde inte läsa PDFConverter-namn")
        return "Ingen"


def create_template_config(
    name: str = "",
    description: str = "",
    counter_start: int = 1,
    validate_vin_flag: bool = True,
    validate_reg_flag: bool = False,
    category: str = "",
    **kwargs: Any,
) -> TemplateConfig:
    """Skapa TemplateConfig med säker normalisering av externa värden."""
    try:
        start = max(1, int(counter_start))
    except (TypeError, ValueError):
        start = 1
    return TemplateConfig(
        name=name,
        description=description,
        category=category,
        counter_prefix=str(kwargs.get("counter_prefix", "") or ""),
        counter_start=start,
        extra_fields=[str(value) for value in kwargs.get("extra_fields", []) or []],
        default_values={
            str(key): str(value)
            for key, value in dict(kwargs.get("default_values", {}) or {}).items()
            if str(key).casefold() not in _NEVER_PREFILL
        },
        pdf_attachments_order=[
            str(value) for value in kwargs.get("pdf_attachments_order", []) or []
        ],
        validate_vin=bool(validate_vin_flag),
        validate_reg=bool(validate_reg_flag),
        output_format=str(kwargs.get("output_format", "pdf") or "pdf"),
        filename_pattern=str(
            kwargs.get(
                "filename_pattern",
                "{number} Intyg {prefix} {vin} {reg} {date}",
            )
        ),
    )


def get_supported_platforms() -> list[str]:
    return ["windows", "darwin", "linux"]


def get_keyboard_shortcuts() -> dict[str, str]:
    return {
        "generate": "Ctrl+G",
        "refresh": "F5",
        "clear": "Ctrl+L",
        "import": "Ctrl+I",
        "open_output": "Ctrl+O",
    }


def get_dependency_status() -> dict[str, Any]:
    """Ge UI och felsökning en snabb, komplett dependency-status."""
    packages = _check_dependencies()
    with _state_lock:
        engine_loaded = _ENGINE_LOADED
        plugin_loaded = _PLUGIN_LOADED
        engine_error = _ENGINE_ERROR
        plugin_error = _PLUGIN_ERROR
    return {
        "packages": packages,
        "all_packages_ok": all(packages.values()),
        "pdf_converter": get_pdf_converter_name(),
        "pdf_available": is_pdf_converter_available(),
        "platform": PlatformHelper.get_platform_name(),
        "python_version": sys.version.split()[0],
        "engine_loaded": engine_loaded,
        "plugin_loaded": plugin_loaded,
        "fully_loaded": engine_loaded and plugin_loaded,
        "engine_error": engine_error,
        "plugin_error": plugin_error,
    }


__all__ = [
    "Plugin",
    "PLUGIN_NAME",
    "PLUGIN_VERSION",
    "PLUGIN_AUTHOR",
    "PLUGIN_PRIORITY",
    "FOLDER_CONFIG",
    "COPYRIGHT",
    "PLUGIN_INFO",
    "PLACEHOLDER_PATTERN",
    "VIN_PATTERN",
    "REG_PATTERN",
    "AUTO_FIELDS",
    "FIELD_LABELS",
    "DEFAULT_VALUES",
    "ConversionMethod",
    "ValidationResult",
    "ScanStatus",
    "TemplateConfig",
    "Template",
    "GenerationResult",
    "FieldValidation",
    "BatchResult",
    "PlatformHelper",
    "VINValidator",
    "RegValidator",
    "PlaceholderParser",
    "PDFConverter",
    "PDFMerger",
    "CounterManager",
    "CertificateGenerator",
    "safe_filename",
    "get_plugin_info",
    "get_version",
    "get_copyright",
    "is_fully_loaded",
    "is_engine_loaded",
    "is_plugin_loaded",
    "get_load_errors",
    "validate_vin",
    "validate_reg",
    "format_vin",
    "find_placeholders",
    "is_auto_field",
    "get_field_label",
    "get_pdf_converter",
    "is_pdf_converter_available",
    "get_pdf_converter_name",
    "create_template_config",
    "get_supported_platforms",
    "get_keyboard_shortcuts",
    "get_dependency_status",
]
